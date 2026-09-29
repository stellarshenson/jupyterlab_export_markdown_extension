"""The `convert` command: a real in-process server, real export handlers.

Each test runs `main()` as the console script would. Nothing is mocked on the
way from the command to the handler, so a flag that stops reaching the API,
or a response the command stops reading, fails here.
"""

import base64
import io
import json
import signal
import sys
from pathlib import Path

import pytest

from jupyterlab_export_markdown_extension import routes
from jupyterlab_export_markdown_extension.cli import main

ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


@pytest.fixture(autouse=True)
def sigint_handler():
    """`convert` sets SIGINT to its default, which its own process drops on
    exit; restore the test runner's handler for the next test."""
    handler = signal.getsignal(signal.SIGINT)
    yield
    signal.signal(signal.SIGINT, handler)


@pytest.fixture
def doc(tmp_path):
    source = tmp_path / "report.md"
    source.write_text(
        "# Quarterly report\n\nBody *text*.\n\n> [!NOTE]\n> Remember this.\n",
        encoding="utf-8",
    )
    return source


def run(capsys, *argv):
    """(exit code, stdout lines, stderr) of one command."""
    rc = main([str(a) for a in argv])
    out, err = capsys.readouterr()
    return rc, out.splitlines(), err


class TestConvert:
    def test_each_format_is_written_beside_the_source(self, capsys, doc):
        rc, out, err = run(capsys, "convert", doc, "--to", "pdf", "docx", "html",
                           "--root", doc.parent)
        assert rc == 0, err
        expected = [doc.with_suffix(f".{fmt}") for fmt in ("pdf", "docx", "html")]
        assert out == [str(p) for p in expected], "stdout must list the written paths only"
        assert expected[0].read_bytes()[:5] == b"%PDF-"
        from docx import Document
        text = "\n".join(p.text for p in Document(io.BytesIO(expected[1].read_bytes())).paragraphs)
        assert "Quarterly report" in text
        assert "<h1" in expected[2].read_text(encoding="utf-8")

    def test_output_path_takes_a_single_format(self, capsys, doc, tmp_path):
        target = tmp_path / "elsewhere.docx"
        rc, out, err = run(capsys, "convert", doc, "--to", "docx", "-o", target,
                           "--root", doc.parent)
        assert rc == 0, err
        assert out == [str(target)]
        assert target.read_bytes()[:2] == b"PK"

    def test_output_path_with_two_formats_is_a_usage_error(self, capsys, doc, tmp_path):
        rc, out, err = run(capsys, "convert", doc, "--to", "pdf", "html",
                           "-o", tmp_path / "x.pdf", "--root", doc.parent)
        assert rc == 2
        assert "-o/--output takes a single --to format" in err
        assert not list(tmp_path.glob("*.pdf")) and not list(tmp_path.glob("*.html"))

    def test_output_onto_the_source_is_refused(self, capsys, tmp_path):
        source = tmp_path / "page.html"
        source.write_text("# Page\n", encoding="utf-8")
        rc, _, err = run(capsys, "convert", source, "--to", "html", "--root", tmp_path)
        assert rc == 2
        assert "is the source file" in err
        assert source.read_text(encoding="utf-8") == "# Page\n"

    def test_a_file_outside_the_root_is_refused(self, capsys, doc, tmp_path):
        root = tmp_path / "root"
        root.mkdir()
        rc, out, err = run(capsys, "convert", doc, "--to", "html", "--root", root)
        assert rc == 2
        assert "--root" in err
        assert out == [] and not doc.with_suffix(".html").exists()

    def test_a_missing_file_is_reported(self, capsys, tmp_path):
        rc, out, err = run(capsys, "convert", tmp_path / "absent.md", "--to", "pdf",
                           "--root", tmp_path)
        assert rc == 2
        assert "does not exist" in err

    def test_images_are_read_only_inside_the_root(self, capsys, tmp_path):
        """The root is the server root: `../img.png` embeds when the root holds
        it and stays a link when it does not - the lab's own boundary."""
        (tmp_path / "img.png").write_bytes(ONE_PIXEL_PNG)
        sub = tmp_path / "sub"
        sub.mkdir()
        source = sub / "doc.md"
        source.write_text("# Doc\n\n![pixel](../img.png)\n", encoding="utf-8")

        rc, _, err = run(capsys, "convert", source, "--to", "html", "--root", tmp_path)
        assert rc == 0, err
        assert "data:image/png;base64" in source.with_suffix(".html").read_text(encoding="utf-8")

        rc, _, err = run(capsys, "convert", source, "--to", "html", "--root", sub)
        assert rc == 0, err
        html = source.with_suffix(".html").read_text(encoding="utf-8")
        assert "data:image/png;base64" not in html
        assert "../img.png" in html
        assert "warning: html: These images were not embedded" in err
        assert "(images: ../img.png)" in err

    def test_html_options_reach_the_export(self, capsys, doc):
        out_html = doc.with_suffix(".html")
        rc, _, err = run(capsys, "convert", doc, "--to", "html", "--root", doc.parent)
        assert rc == 0, err
        default = out_html.read_text(encoding="utf-8")
        assert "font-size: 12pt" in default and "color-scheme: light" in default
        assert "NOTE:" not in default

        rc, _, err = run(capsys, "convert", doc, "--to", "html", "--root", doc.parent,
                         "--theme", "dark", "--font-size", "small", "--alert-labels")
        assert rc == 0, err
        chosen = out_html.read_text(encoding="utf-8")
        assert "font-size: 10pt" in chosen, "--font-size did not reach exportFontSize"
        assert "color-scheme: dark" in chosen, "--theme did not reach htmlTheme"
        assert "NOTE:" in chosen, "--alert-labels did not reach showAlertLabels"

    def test_theme_reaches_the_docx_renderer(self, capsys, doc, monkeypatch):
        """DOCX uses the theme only to rasterise SVG and Mermaid, so assert on
        the value the handler resolves rather than on image pixels."""
        seen = []
        original = routes.ExportHandlerBase.color_scheme_for

        def record(theme):
            seen.append(theme)
            return original(theme)

        monkeypatch.setattr(routes.ExportHandlerBase, "color_scheme_for", staticmethod(record))
        rc, _, err = run(capsys, "convert", doc, "--to", "docx", "--root", doc.parent,
                         "--theme", "dark")
        assert rc == 0, err
        assert seen == ["dark"], f"docxTheme sent as {seen}"

    def test_a_failed_format_is_reported_and_the_others_still_written(
        self, capsys, doc, monkeypatch
    ):
        def fail(*_args, **_kwargs):
            raise RuntimeError("reportlab exploded")

        monkeypatch.setattr(routes.ExportHandlerBase, "convert_docx_to_pdf", fail)
        rc, out, err = run(capsys, "convert", doc, "--to", "pdf", "html",
                           "--root", doc.parent)
        assert rc == 1
        assert "error: pdf: reportlab exploded" in err
        assert out == [str(doc.with_suffix(".html"))]
        assert not doc.with_suffix(".pdf").exists()

    def test_export_warnings_go_to_stderr_without_failing(self, capsys, tmp_path):
        source = tmp_path / "mix.md"
        source.write_text("# Mix\n\n```mermaid\nnot a diagram at all {{{\n```\n",
                          encoding="utf-8")
        rc, out, err = run(capsys, "convert", source, "--to", "docx", "--root", tmp_path)
        assert rc == 0, err
        assert out == [str(source.with_suffix(".docx"))]
        assert "warning: docx: Mermaid could not draw these diagrams" in err
        assert "(mermaid blocks 1)" in err, "positions must count from 1"

    def test_a_repeated_to_keeps_every_format(self, capsys, doc):
        rc, out, err = run(capsys, "convert", doc, "--to", "html", "--to", "docx",
                           "--root", doc.parent)
        assert rc == 0, err
        assert out == [str(doc.with_suffix(".html")), str(doc.with_suffix(".docx"))]

    def test_the_users_jupyter_config_is_not_read(self, capsys, tmp_path, monkeypatch):
        """A contents root in the user's server config made the handlers read
        another file of the same name and exit 0."""
        root, other, cfg = tmp_path / "root", tmp_path / "other", tmp_path / "cfg"
        for folder in (root, other, cfg):
            folder.mkdir()
        (root / "doc.md").write_text("# Right file\n", encoding="utf-8")
        (other / "doc.md").write_text("# Wrong file\n", encoding="utf-8")
        (cfg / "jupyter_server_config.json").write_text(
            json.dumps({"FileContentsManager": {"root_dir": str(other)}}), encoding="utf-8")
        monkeypatch.setenv("JUPYTER_CONFIG_DIR", str(cfg))
        rc, _, err = run(capsys, "convert", root / "doc.md", "--to", "html", "--root", root)
        assert rc == 0, err
        html = (root / "doc.html").read_text(encoding="utf-8")
        assert "Right file" in html and "Wrong file" not in html

    def test_a_transport_error_fails_its_format_only(self, capsys, doc, monkeypatch):
        """raise_error=False does not cover a connection closed early - the
        response-size cap ends that way - so it must be caught per format."""
        from tornado.httpclient import AsyncHTTPClient
        from tornado.simple_httpclient import HTTPStreamClosedError
        real_fetch = AsyncHTTPClient.fetch

        def fetch(self, url, **kwargs):
            if url.endswith("/export/pdf"):
                raise HTTPStreamClosedError("Stream closed")
            return real_fetch(self, url, **kwargs)

        monkeypatch.setattr(AsyncHTTPClient, "fetch", fetch)
        rc, out, err = run(capsys, "convert", doc, "--to", "pdf", "html",
                           "--root", doc.parent)
        assert rc == 1
        assert "error: pdf: Stream closed" in err
        assert out == [str(doc.with_suffix(".html"))]

    def test_a_symlinked_file_exports_from_where_it_is_linked(self, capsys, tmp_path):
        """The server joins root and path without resolving links, as the lab
        does: images resolve from the link's folder, the output lands beside
        the link, and a target outside the root is no reason to refuse."""
        root, real = tmp_path / "p", tmp_path / "real"
        (root / "sub").mkdir(parents=True)
        real.mkdir()
        (real / "r.md").write_text("# R\n\n![i](sub/pix.png)\n", encoding="utf-8")
        (root / "sub" / "pix.png").write_bytes(ONE_PIXEL_PNG)
        (root / "link.md").symlink_to(real / "r.md")
        rc, out, err = run(capsys, "convert", root / "link.md", "--to", "html",
                           "--root", root)
        assert rc == 0, err
        assert out == [str(root / "link.html")]
        assert "data:image/png;base64" in (root / "link.html").read_text(encoding="utf-8")

    def test_a_folder_reached_through_a_link_is_inside_the_root(
        self, capsys, tmp_path, monkeypatch
    ):
        """The current directory arrives resolved and a typed path keeps its
        links; the two must still compare, or a file in the root is refused."""
        real = tmp_path / "real"
        real.mkdir()
        (real / "doc.md").write_text("# Doc\n", encoding="utf-8")
        (tmp_path / "link").symlink_to(real)
        monkeypatch.chdir(tmp_path / "link")
        rc, out, err = run(capsys, "convert", tmp_path / "link" / "doc.md", "--to", "html")
        assert rc == 0, err
        assert (real / "doc.html").exists()

    def test_image_syntax_inside_code_is_left_as_written(self, capsys, tmp_path):
        """A code sample shows image syntax; it places no image, so it is
        neither rewritten to base64 nor reported as missing."""
        (tmp_path / "logo.png").write_bytes(ONE_PIXEL_PNG)
        source = tmp_path / "readme.md"
        source.write_text(
            "# Readme\n\n```markdown\n![logo](logo.png)\n```\n\n"
            "Write `![x](missing.png)` for an image.\n\n<!-- ![old](gone.png) -->\n",
            encoding="utf-8")
        rc, _, err = run(capsys, "convert", source, "--to", "html", "--root", tmp_path)
        assert rc == 0, err
        html = source.with_suffix(".html").read_text(encoding="utf-8")
        assert "data:image/png" not in html, "a code sample was rewritten to base64"
        assert "warning:" not in err

    def test_a_refused_image_stays_out_of_docx(self, capsys, tmp_path, monkeypatch):
        """htmldocx loaded any path left in an <img> from the working
        directory: a missing sub/chart.png came out as the root's chart.png."""
        (tmp_path / "chart.png").write_bytes(ONE_PIXEL_PNG)
        (tmp_path / "sub").mkdir()
        source = tmp_path / "sub" / "doc.md"
        source.write_text("# Doc\n\n![c](chart.png)\n", encoding="utf-8")
        monkeypatch.chdir(tmp_path)
        rc, _, err = run(capsys, "convert", source, "--to", "docx", "--root", tmp_path)
        assert rc == 0, err
        from docx import Document
        document = Document(str(source.with_suffix(".docx")))
        assert not document.inline_shapes, "an image the root boundary refused was embedded"
        assert "<image: chart.png>" in "\n".join(p.text for p in document.paragraphs)
        assert "(images: chart.png)" in err

    def test_a_refused_download_is_not_fetched_for_docx(self, capsys, tmp_path):
        """The embed pass refuses a private host (SSRF guard); htmldocx fetched
        it anyway with a plain urlopen."""
        import http.server
        import threading

        served = []

        class Pixel(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                served.append(self.path)
                self.send_response(200)
                self.send_header("Content-Type", "image/png")
                self.end_headers()
                self.wfile.write(ONE_PIXEL_PNG)

            def log_message(self, *args):
                pass

        server = http.server.HTTPServer(("127.0.0.1", 0), Pixel)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        try:
            url = f"http://127.0.0.1:{server.server_port}/p.png"
            source = tmp_path / "doc.md"
            source.write_text(f"# Doc\n\n![p]({url})\n", encoding="utf-8")
            rc, _, err = run(capsys, "convert", source, "--to", "html", "docx",
                             "--root", tmp_path)
        finally:
            server.shutdown()
        assert rc == 0, err
        assert served == [], f"the private host was fetched: {served}"
        from docx import Document
        assert not Document(str(source.with_suffix(".docx"))).inline_shapes
        for fmt in ("html", "docx"):
            assert f"warning: {fmt}: These images were not embedded" in err
        assert err.count(url) == 2

    def test_an_ignored_sigint_stays_ignored(self, capsys, doc):
        """A job started with SIGINT ignored keeps it ignored - the rule
        asyncio itself follows."""
        signal.signal(signal.SIGINT, signal.SIG_IGN)
        rc, _, err = run(capsys, "convert", doc, "--to", "html", "--root", doc.parent)
        assert rc == 0, err
        assert signal.getsignal(signal.SIGINT) is signal.SIG_IGN

    def test_an_error_without_text_is_named(self, capsys, doc, monkeypatch):
        """python-docx raises UnrecognizedImageError with no message; the
        error line said only 'HTTP 500: Internal Server Error'."""
        def fail(*_args, **_kwargs):
            raise RuntimeError()

        monkeypatch.setattr(routes.ExportHandlerBase, "convert_docx_to_pdf", fail)
        rc, _, err = run(capsys, "convert", doc, "--to", "pdf", "--root", doc.parent)
        assert rc == 1
        assert "error: pdf: RuntimeError" in err

    def test_check_names_the_remedy_first(self, capsys, monkeypatch):
        from jupyterlab_export_markdown_extension import cli
        monkeypatch.setattr(cli, "_try_launch_chromium",
                            lambda: (False, "Executable doesn't exist\n  playwright install"))
        rc, _, err = run(capsys, "check")
        assert rc == 1
        assert err.startswith(
            "FAIL: Chromium cannot launch. Fix with:  jupyterlab-export-markdown-extension install")

    def test_an_output_linked_to_the_source_is_refused(self, capsys, tmp_path):
        source = tmp_path / "s.md"
        source.write_text("# S\n", encoding="utf-8")
        (tmp_path / "s.html").symlink_to(source)
        rc, _, err = run(capsys, "convert", source, "--to", "html", "--root", tmp_path)
        assert rc == 2
        assert "is the source file" in err
        assert source.read_text(encoding="utf-8") == "# S\n"


class TestHelp:
    @pytest.mark.parametrize("argv", [["--help"], ["convert", "--help"]])
    def test_help_carries_worked_examples(self, capsys, argv):
        with pytest.raises(SystemExit) as exit_info:
            main(argv)
        assert exit_info.value.code == 0
        assert "examples:" in capsys.readouterr().out


def test_the_installed_skill_is_the_repository_copy():
    """pyproject.toml maps the agent skill into the wheel as shared-data; this
    fails when the mapping is lost or the environment holds an older wheel."""
    skill = Path("skills/jupyterlab-export-markdown-extension/SKILL.md")
    repository = Path(__file__).parents[2] / ".agents" / skill
    installed = Path(sys.prefix) / "share/jupyter/agents" / skill
    assert installed.read_text(encoding="utf-8") == repository.read_text(encoding="utf-8")
